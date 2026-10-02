"""CRM 주소록·안건 (README 32 · /api/crm). 실 Postgres 위에서 자원별 정상 경로·권한 거부(403+감사)·검증 422·409."""

import io
from collections.abc import Callable
from datetime import UTC, datetime, timedelta

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import AuditEvent, CrmImport, CrmOrgOwner, User
from tests.conftest import PW_HASH, Users, csrf, login

API = "/api/crm"


async def _as(client: AsyncClient, email: str) -> dict[str, str]:
    r = await login(client, email)
    assert r.status_code == 204, r.text
    return csrf(client)


async def _audits(db: AsyncSession, kind: str) -> list[AuditEvent]:
    rows = await db.execute(select(AuditEvent).where(AuditEvent.kind == kind))
    return list(rows.scalars())


def _ago(days: float) -> str:
    return (datetime.now(UTC) - timedelta(days=days)).isoformat()


async def _org(client: AsyncClient, h: dict, name: str, org_type: str = "pension", **kw) -> dict:
    r = await client.post(f"{API}/orgs", json={"name": name, "org_type": org_type, **kw}, headers=h)
    assert r.status_code == 201, r.text
    return r.json()["org"]


async def _contact(client: AsyncClient, h: dict, org_id: str, name: str, **kw) -> dict:
    r = await client.post(f"{API}/contacts", json={"org_id": org_id, "name": name, **kw}, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


async def _deal(client: AsyncClient, h: dict, org_id: str, **kw) -> dict:
    body = {"kind": "fundraising", "name": "인프라 2호 펀드 출자", "org_id": org_id, **kw}
    r = await client.post(f"{API}/deals", json=body, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


async def _activity(client: AsyncClient, h: dict, org_id: str, **kw) -> dict:
    body = {
        "kind": "meeting",
        "occurred_at": _ago(1),
        "org_id": org_id,
        "body": "제안서 설명",
        **kw,
    }
    r = await client.post(f"{API}/activities", json=body, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


async def test_requires_login(client: AsyncClient, users: Users) -> None:
    r = await client.get(f"{API}/orgs")
    assert r.status_code == 401
    assert r.json()["code"] == "AUTH_REQUIRED"


async def test_orgs_create_list_filters(client: AsyncClient, users: Users, app_db) -> None:
    """정규화 이름 중복 경고(막지 않음)·q·type·owner·stale_days·정렬·contact_count·last_contacted_at."""
    h = await _as(client, "sujin@zenith.co.kr")
    first = await _org(
        client,
        h,
        " (주)미래인프라 자산운용 ",
        "am",
        memo="공동투자 상대.",
        owner_ids=[str(users.lead.id)],
    )
    assert first["name"] == "(주)미래인프라 자산운용"
    assert first["owners"] == [{"id": str(users.lead.id), "name": "박동현"}]
    assert (first["contact_count"], first["last_contacted_at"], first["contacts"]) == (0, None, [])
    assert first["memo"] == "공동투자 상대."

    # 정규화 이름이 같으면 만들되 duplicates 로 알린다
    r = await client.post(
        f"{API}/orgs", json={"name": "미래인프라자산운용㈜", "org_type": "am"}, headers=h
    )
    assert r.status_code == 201
    assert r.json()["duplicates"] == [{"id": first["id"], "name": first["name"]}]
    dup = r.json()["org"]

    hangyeol = await _org(client, h, "한결연금공단", owner_ids=[str(users.member.id)])
    await _org(client, h, "누리공제회", "mutual_aid")

    async def names(**params) -> list[str]:
        r = await client.get(f"{API}/orgs", params=params)
        assert r.status_code == 200, r.text
        assert r.json()["next_cursor"] is None
        return [o["name"] for o in r.json()["items"]]

    assert await names() == [
        "누리공제회",
        "(주)미래인프라 자산운용",
        "미래인프라자산운용㈜",
        "한결연금공단",
    ]
    assert await names(q="주식회사 미래 인프라") == [
        "(주)미래인프라 자산운용",
        "미래인프라자산운용㈜",
    ]
    assert await names(type="mutual_aid") == ["누리공제회"]
    assert await names(type="nope") == []
    assert await names(owner=str(users.member.id)) == ["한결연금공단"]
    assert await names(owner="not-a-uuid") == []

    # 담당자 수는 개인정보를 지운 사람을 빼고, 마지막 연락은 지우지 않은 활동의 최댓값
    c1 = await _contact(client, h, hangyeol["id"], "박서준")
    await _contact(client, h, hangyeol["id"], "이하은")
    await _activity(client, h, hangyeol["id"], occurred_at=_ago(40))
    recent = await _activity(client, h, dup["id"], occurred_at=_ago(3))
    hx = await _as(client, "jaeho@zenith.co.kr")  # 같은 클라이언트를 임원으로 갈아탄다
    assert (await client.post(f"{API}/contacts/{c1['id']}/erase", headers=hx)).status_code == 200
    items = {o["name"]: o for o in (await client.get(f"{API}/orgs")).json()["items"]}
    assert items["한결연금공단"]["contact_count"] == 1
    assert items["누리공제회"]["last_contacted_at"] is None
    # 30일 넘게 연락 없음 = 40일 전 활동뿐인 기관 + 활동이 없는 기관
    assert await names(stale_days=30) == ["누리공제회", "(주)미래인프라 자산운용", "한결연금공단"]
    assert await names(stale_days=0) == await names()
    # 지운 활동은 마지막 연락에서 빠진다
    assert (await client.delete(f"{API}/activities/{recent['id']}", headers=hx)).status_code == 204
    assert "미래인프라자산운용㈜" in await names(stale_days=1)


async def test_org_validation_and_patch(client: AsyncClient, users: Users, app_db) -> None:
    h = await _as(client, "sujin@zenith.co.kr")
    for bad in (
        {"name": "  ", "org_type": "bank"},
        {"name": "가" * 101, "org_type": "bank"},
        {"name": "세움은행", "org_type": "nope"},
        {"name": "세움은행"},
        {"name": "세움은행", "org_type": "bank", "memo": "x" * 2001},
        {"name": "세움은행", "org_type": "bank", "owner_ids": ["not-a-uuid"]},
    ):
        r = await client.post(f"{API}/orgs", json=bad, headers=h)
        assert r.status_code == 422, bad
        assert r.json()["code"] == "VALIDATION"

    retired = User(name="퇴사자", email="gone@zenith.co.kr", pw_hash=PW_HASH, active=False)
    app_db.add(retired)
    await app_db.flush()
    r = await client.post(
        f"{API}/orgs",
        json={"name": "세움은행", "org_type": "bank", "owner_ids": [str(retired.id)]},
        headers=h,
    )
    assert r.status_code == 422
    org = await _org(client, h, "세움은행", "bank")
    # 퇴사자가 이미 앉아 있으면 그대로 두고 다른 칸을 고칠 수 있다
    app_db.add(CrmOrgOwner(org_id=org["id"], user_id=retired.id))
    await app_db.flush()
    url = f"{API}/orgs/{org['id']}"
    r = await client.patch(
        url, json={"memo": "메모", "owner_ids": [str(retired.id), str(users.lead.id)]}, headers=h
    )
    assert r.status_code == 200, r.text
    assert {p["name"] for p in r.json()["owners"]} == {"퇴사자", "박동현"}
    r = await client.patch(url, json={"owner_ids": [str(users.lead.id)]}, headers=h)
    assert [p["name"] for p in r.json()["owners"]] == ["박동현"]
    r = await client.patch(url, json={"name": "㈜세움 은행", "org_type": "etc"}, headers=h)
    assert (r.json()["name"], r.json()["org_type"], r.json()["memo"]) == (
        "㈜세움 은행",
        "etc",
        "메모",
    )
    assert [
        o["name"] for o in (await client.get(f"{API}/orgs", params={"q": "세움은"})).json()["items"]
    ] == ["㈜세움 은행"]
    assert (await client.patch(url, json={"name": None}, headers=h)).status_code == 422
    assert (
        await client.patch(url, json={"owner_ids": [str(retired.id)]}, headers=h)
    ).status_code == 422
    missing = f"{API}/orgs/00000000-0000-0000-0000-000000000000"
    assert (await client.get(missing)).status_code == 404
    assert (await client.patch(missing, json={"memo": "x"}, headers=h)).status_code == 404


async def test_org_delete(
    client: AsyncClient, make_client: Callable[[], AsyncClient], users: Users, app_db
) -> None:
    """exec 만, 사유 필수, 진행 중 안건 409. soft delete 뒤 상세·안건·활동 404, 본문 참조 422, 내보내기 제외."""
    h = await _as(client, "sujin@zenith.co.kr")
    org = await _org(client, h, "다온생명보험", "insurance")
    deal = await _deal(client, h, org["id"])
    await _activity(client, h, org["id"], deal_id=deal["id"])
    await _contact(client, h, org["id"], "정유나")
    url = f"{API}/orgs/{org['id']}"

    r = await client.request("DELETE", url, json={"reason": "중복"}, headers=h)
    assert r.status_code == 403
    denied = await _audits(app_db, "acl.denied")
    assert [(a.actor_id, str(a.target_id), a.detail["reason"]) for a in denied] == [
        (users.member.id, org["id"], "org_delete")
    ]

    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        assert (
            await boss.request("DELETE", url, json={"reason": " "}, headers=hx)
        ).status_code == 422
        assert (await boss.request("DELETE", url, headers=hx)).status_code == 422
        r = await boss.request("DELETE", url, json={"reason": "중복 등록"}, headers=hx)
        assert r.status_code == 409
        assert "1건" in r.json()["detail"]
        dropped = await boss.post(
            f"{API}/deals/{deal['id']}/drop", json={"reason": "한도 소진", "version": 1}, headers=hx
        )
        assert dropped.status_code == 200
        r = await boss.request("DELETE", url, json={"reason": " 중복 등록 "}, headers=hx)
        assert r.status_code == 204, r.text
        assert (
            await boss.request("DELETE", url, json={"reason": "x"}, headers=hx)
        ).status_code == 404
        exported = await boss.post(f"{API}/export", params={"entity": "contacts"}, headers=hx)
        assert load_workbook(io.BytesIO(exported.content)).active.max_row == 1  # 머리글만

    [ev] = await _audits(app_db, "crm.org_deleted")
    assert (str(ev.target_id), ev.detail) == (
        org["id"],
        {"name": "다온생명보험", "reason": "중복 등록"},
    )
    assert (await client.get(url)).status_code == 404
    assert (await client.get(f"{API}/orgs")).json()["items"] == []
    assert (await client.get(f"{API}/deals")).json()["items"] == []
    assert (await client.get(f"{API}/deals/{deal['id']}")).status_code == 404
    reopen = await client.post(f"{API}/deals/{deal['id']}/reopen", json={"version": 2}, headers=h)
    assert reopen.status_code == 404
    assert (await client.get(f"{API}/activities", params={"org_id": org["id"]})).status_code == 404
    assert (
        await client.get(f"{API}/activities", params={"deal_id": deal["id"]})
    ).status_code == 404
    for path, body in (
        ("contacts", {"org_id": org["id"], "name": "새 사람"}),
        ("deals", {"kind": "investment", "name": "새 안건", "org_id": org["id"]}),
        ("activities", {"kind": "call", "occurred_at": _ago(0), "org_id": org["id"], "body": "x"}),
    ):
        r = await client.post(f"{API}/{path}", json=body, headers=h)
        assert r.status_code == 422, path


async def test_contacts_move_and_erase(
    client: AsyncClient, make_client: Callable[[], AsyncClient], users: Users, app_db
) -> None:
    h = await _as(client, "sujin@zenith.co.kr")
    daon = await _org(client, h, "다온생명보험", "insurance")
    mirae = await _org(client, h, "미래인프라자산운용", "am")
    c = await _contact(
        client, h, daon["id"], " 한도윤 ", title="인프라본부 상무", email="dy.han@mirae.example"
    )
    assert (c["name"], c["erased"], c["moves"], c["phone"]) == ("한도윤", False, [], "")
    for bad in (
        {"org_id": daon["id"], "name": " "},
        {"org_id": daon["id"], "name": "x", "email": "no-at-sign"},
        {"org_id": daon["id"], "name": "x", "memo": "x" * 2001},
        {"org_id": daon["id"], "name": "x" * 201},
        {"org_id": "00000000-0000-0000-0000-000000000000", "name": "x"},
    ):
        assert (await client.post(f"{API}/contacts", json=bad, headers=h)).status_code == 422, bad

    url = f"{API}/contacts/{c['id']}"
    act = await _activity(client, h, daon["id"], contact_ids=[c["id"]])
    r = await client.patch(url, json={"org_id": mirae["id"], "phone": "02-0000-4001"}, headers=h)
    assert r.status_code == 200, r.text
    moved = r.json()
    assert (moved["org_id"], moved["phone"]) == (mirae["id"], "02-0000-4001")
    assert [m["from"] for m in moved["moves"]] == [{"id": daon["id"], "name": "다온생명보험"}]
    # 같은 기관을 다시 보내면 이직이 아니다
    again = await client.patch(url, json={"org_id": mirae["id"]}, headers=h)
    assert len(again.json()["moves"]) == 1
    bad_move = await client.patch(
        url, json={"org_id": "00000000-0000-0000-0000-000000000000"}, headers=h
    )
    assert bad_move.status_code == 422
    # 옛 활동은 옛 기관에 남는다
    old = (await client.get(f"{API}/activities", params={"org_id": daon["id"]})).json()["items"]
    assert [a["id"] for a in old] == [act["id"]]
    assert (await client.get(f"{API}/orgs/{mirae['id']}")).json()["contacts"][0]["name"] == "한도윤"

    # 개인정보 삭제 — exec 만(403 이 404 보다 먼저)
    r = await client.post(f"{url}/erase", headers=h)
    assert r.status_code == 403
    r = await client.post(f"{API}/contacts/00000000-0000-0000-0000-000000000000/erase", headers=h)
    assert r.status_code == 403
    assert len(await _audits(app_db, "acl.denied")) == 2
    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        r = await boss.post(f"{url}/erase", headers=hx)
        assert r.status_code == 200
        erased = r.json()
        assert {k: erased[k] for k in ("name", "title", "phone", "email", "memo", "erased")} == {
            "name": "",
            "title": "",
            "phone": "",
            "email": "",
            "memo": "",
            "erased": True,
        }
        assert (erased["org_id"], len(erased["moves"])) == (mirae["id"], 1)
        assert (await boss.post(f"{url}/erase", headers=hx)).status_code == 200  # 멱등
        missing = f"{API}/contacts/00000000-0000-0000-0000-000000000000/erase"
        assert (await boss.post(missing, headers=hx)).status_code == 404
    events = await _audits(app_db, "crm.contact_erased")
    assert [(str(e.target_id), e.detail) for e in events] == [(c["id"], None)] * 2

    r = await client.patch(url, json={"memo": "다시"}, headers=h)
    assert r.status_code == 409
    assert r.json()["code"] == "CONFLICT"
    detail = (await client.get(f"{API}/orgs/{mirae['id']}")).json()
    assert (detail["contact_count"], detail["contacts"][0]["erased"]) == (0, True)  # 상세엔 남는다
    # 활동 참석자 이름은 "" 로
    old = (await client.get(f"{API}/activities", params={"org_id": daon["id"]})).json()["items"]
    assert old[0]["contacts"] == [{"id": c["id"], "name": ""}]


async def test_deal_lifecycle(
    client: AsyncClient, make_client: Callable[[], AsyncClient], users: Users, app_db
) -> None:
    """version 잠금(누락·어긋남 409)·단계 이동·성사(won)·성사 되돌리기(exec 만)·중단·재개."""
    h = await _as(client, "sujin@zenith.co.kr")
    org = await _org(client, h, "한결연금공단")
    deal = await _deal(
        client,
        h,
        org["id"],
        owner_id=str(users.lead.id),
        amount_krw=50_000_000_000,
        target_date="2026-12-01",
        next_action=" 실사 질의서 답변 송부 ",
        next_action_due="2026-10-06",
    )
    assert (deal["stage"], deal["status"], deal["version"]) == ("contact", "open", 1)
    assert (deal["owner"]["name"], deal["org"]["name"]) == ("박동현", "한결연금공단")
    assert (deal["amount_krw"], deal["next_action"], deal["closed_at"]) == (
        50_000_000_000,
        "실사 질의서 답변 송부",
        None,
    )
    assert [(e["from_stage"], e["to_stage"], e["action"]) for e in deal["events"]] == [
        (None, "contact", "move")
    ]
    assert deal["events"][0]["actor"]["name"] == "김수진"
    for bad in (
        {"kind": "nope", "name": "x", "org_id": org["id"]},
        {"kind": "investment", "name": " ", "org_id": org["id"]},
        {"kind": "investment", "name": "x", "org_id": org["id"], "amount_krw": -1},
        {"kind": "investment", "name": "x", "org_id": org["id"], "amount_krw": 1.5},
        {"kind": "investment", "name": "x", "org_id": org["id"], "target_date": "2026-02-30"},
        {"kind": "investment", "name": "x", "org_id": org["id"], "next_action_due": "2026-10-01"},
        {
            "kind": "investment",
            "name": "x",
            "org_id": org["id"],
            "owner_id": "00000000-0000-0000-0000-000000000000",
        },
    ):
        assert (await client.post(f"{API}/deals", json=bad, headers=h)).status_code == 422, bad

    url = f"{API}/deals/{deal['id']}"
    assert (
        await client.patch(url, json={"name": "새 이름"}, headers=h)
    ).status_code == 409  # version 누락
    r = await client.patch(url, json={"version": 9, "name": "새 이름"}, headers=h)
    assert (r.status_code, r.json()["code"]) == (409, "CONFLICT")
    assert (
        await client.patch(url, json={"version": 1, "org_id": org["id"]}, headers=h)
    ).status_code == 422
    assert (
        await client.patch(url, json={"version": 1, "stage": "closed"}, headers=h)
    ).status_code == 422  # 다른 종류의 단계
    # 할 일을 비우면서 기한이 남으면 422 — 단계만 옮기는 요청은 안 본다
    assert (
        await client.patch(url, json={"version": 1, "next_action": ""}, headers=h)
    ).status_code == 422
    r = await client.patch(url, json={"version": 1, "stage": "diligence"}, headers=h)
    assert r.status_code == 200, r.text
    assert (r.json()["stage"], r.json()["version"], r.json()["status"]) == ("diligence", 2, "open")
    r = await client.patch(
        url, json={"version": 2}, headers=h
    )  # 칸이 안 바뀌어도 version 은 오른다
    assert r.json()["version"] == 3
    r = await client.patch(
        url,
        json={
            "version": 3,
            "owner_id": None,
            "amount_krw": None,
            "next_action": "",
            "next_action_due": None,
        },
        headers=h,
    )
    assert (r.json()["owner"], r.json()["amount_krw"], r.json()["next_action_due"]) == (
        None,
        None,
        None,
    )

    # 마지막 단계 = 성사. 사원은 되돌릴 수 없다
    r = await client.patch(url, json={"version": 4, "stage": "committed"}, headers=h)
    won = r.json()
    assert (won["status"], won["stage"], won["closed_at"] is not None) == ("won", "committed", True)
    r = await client.patch(url, json={"version": 5, "stage": "committee"}, headers=h)
    assert (r.status_code, r.json()["detail"]) == (409, "끝난 안건은 단계를 옮길 수 없습니다.")
    r = await client.patch(url, json={"version": 5, "name": "성사 뒤 이름 고치기"}, headers=h)
    assert r.status_code == 200  # 단계 외 칸은 고칠 수 있다
    assert (
        await client.post(f"{url}/drop", json={"reason": "x", "version": 6}, headers=h)
    ).status_code == 409
    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        r = await boss.patch(url, json={"version": 6, "stage": "committee"}, headers=hx)
        assert r.status_code == 200, r.text
        undone = r.json()
        assert (undone["status"], undone["closed_at"], undone["version"]) == ("open", None, 7)
        assert undone["events"][0]["actor"]["name"] == "최재호"
    assert [(e["from_stage"], e["to_stage"]) for e in undone["events"]] == [
        ("committed", "committee"),
        ("diligence", "committed"),
        ("contact", "diligence"),
        (None, "contact"),
    ]

    # 중단 — 사유 422 → version 409 → 상태 409
    assert (
        await client.post(f"{url}/drop", json={"reason": " ", "version": 7}, headers=h)
    ).status_code == 422
    assert (
        await client.post(f"{url}/drop", json={"reason": "한도 소진"}, headers=h)
    ).status_code == 409
    r = await client.post(f"{url}/drop", json={"reason": " 한도 소진 ", "version": 7}, headers=h)
    dropped = r.json()
    assert (dropped["status"], dropped["drop_reason"], dropped["stage"], dropped["version"]) == (
        "dropped",
        "한도 소진",
        "committee",
        8,
    )
    assert dropped["closed_at"] is not None
    assert (dropped["events"][0]["action"], dropped["events"][0]["reason"]) == ("drop", "한도 소진")
    r = await client.post(f"{url}/drop", json={"reason": "또", "version": 8}, headers=h)
    assert r.json()["detail"] == "진행 중인 안건만 중단할 수 있습니다."
    async with make_client() as boss:  # 중단된 안건은 임원도 단계를 못 옮긴다
        hx = await _as(boss, "jaeho@zenith.co.kr")
        assert (
            await boss.patch(url, json={"version": 8, "stage": "contact"}, headers=hx)
        ).status_code == 409
    assert (await client.post(f"{url}/reopen", json={"version": 7}, headers=h)).status_code == 409
    r = await client.post(f"{url}/reopen", json={"version": 8}, headers=h)
    reopened = r.json()
    assert (
        reopened["status"],
        reopened["drop_reason"],
        reopened["closed_at"],
        reopened["version"],
    ) == ("open", None, None, 9)
    assert reopened["events"][0]["action"] == "reopen"
    r = await client.post(f"{url}/reopen", json={"version": 9}, headers=h)
    assert r.json()["detail"] == "중단된 안건만 재개할 수 있습니다."

    # 목록 필터 — 등록 최신순
    other = await _deal(client, h, org["id"], kind="disposal", name="메디컬센터 매각")
    listed = (await client.get(f"{API}/deals")).json()["items"]
    assert [d["id"] for d in listed] == [other["id"], deal["id"]]
    assert "events" not in listed[0]
    assert [
        d["id"]
        for d in (await client.get(f"{API}/deals", params={"kind": "disposal"})).json()["items"]
    ] == [other["id"]]
    assert (await client.get(f"{API}/deals", params={"status": "won"})).json()["items"] == []
    assert (await client.get(f"{API}/deals", params={"org_id": "x"})).json()["items"] == []


async def test_activities(
    client: AsyncClient, make_client: Callable[[], AsyncClient], users: Users, app_db
) -> None:
    """교차 검증(담당자·안건이 그 기관 것)·미래 일시 422·can_edit·작성자/exec 만 고치고 지움(403+감사)."""
    h = await _as(client, "sujin@zenith.co.kr")
    org = await _org(client, h, "한결연금공단")
    other = await _org(client, h, "누리공제회", "mutual_aid")
    c1 = await _contact(client, h, org["id"], "박서준")
    c2 = await _contact(client, h, org["id"], "이하은")
    outsider = await _contact(client, h, other["id"], "최민호")
    deal = await _deal(client, h, org["id"])
    other_deal = await _deal(client, h, other["id"])

    act = await _activity(
        client,
        h,
        org["id"],
        deal_id=deal["id"],
        contact_ids=[c2["id"], c1["id"], c1["id"]],
        user_ids=[str(users.lead.id), str(users.member.id)],
        body="  2호 펀드 제안서 설명.  ",
    )
    assert act["body"] == "2호 펀드 제안서 설명."
    assert [p["name"] for p in act["contacts"]] == ["박서준", "이하은"]
    assert [p["name"] for p in act["users"]] == ["김수진", "박동현"]
    assert (act["org"]["name"], act["deal"]["id"], act["author"]["name"], act["can_edit"]) == (
        "한결연금공단",
        deal["id"],
        "김수진",
        True,
    )
    plain = await _activity(
        client, h, org["id"], kind="call", occurred_at=_ago(20), body="일정 조율"
    )
    assert plain["deal"] is None

    base = {"kind": "call", "occurred_at": _ago(0), "org_id": org["id"], "body": "x"}
    for bad in (
        {**base, "occurred_at": (datetime.now(UTC) + timedelta(minutes=10)).isoformat()},
        {**base, "occurred_at": "2026-09-01T10:00:00"},  # 시간대 없음
        {**base, "kind": "visit"},
        {**base, "body": " "},
        {**base, "body": "x" * 5001},
        {**base, "contact_ids": [outsider["id"]]},
        {**base, "deal_id": other_deal["id"]},
        {**base, "deal_id": "00000000-0000-0000-0000-000000000000"},
        {**base, "user_ids": ["00000000-0000-0000-0000-000000000000"]},
        {**base, "org_id": "00000000-0000-0000-0000-000000000000"},
    ):
        r = await client.post(f"{API}/activities", json=bad, headers=h)
        assert r.status_code == 422, bad
    soon = (datetime.now(UTC) + timedelta(minutes=2)).isoformat()  # 시계 오차 5분은 봐준다
    await _activity(client, h, other["id"], occurred_at=soon)

    r = await client.get(f"{API}/activities")
    assert (r.status_code, r.json()["detail"]) == (422, "org_id 나 deal_id 가 필요합니다.")
    by_org = (await client.get(f"{API}/activities", params={"org_id": org["id"]})).json()["items"]
    assert [a["id"] for a in by_org] == [act["id"], plain["id"]]  # 안건에 건 기록도 기관 타임라인에
    by_deal = (await client.get(f"{API}/activities", params={"deal_id": deal["id"]})).json()[
        "items"
    ]
    assert [a["id"] for a in by_deal] == [act["id"]]
    missing = {"org_id": "00000000-0000-0000-0000-000000000000"}
    assert (await client.get(f"{API}/activities", params=missing)).status_code == 404

    url = f"{API}/activities/{act['id']}"
    async with make_client() as lead:
        hl = await _as(lead, "donghyun@zenith.co.kr")
        seen = (await lead.get(f"{API}/activities", params={"org_id": org["id"]})).json()["items"]
        assert [a["can_edit"] for a in seen] == [False, False]
        assert (await lead.patch(url, json={"body": "남의 글"}, headers=hl)).status_code == 403
        assert (await lead.delete(url, headers=hl)).status_code == 403
    denied = await _audits(app_db, "acl.denied")
    assert sorted(a.detail["reason"] for a in denied) == ["activity_delete", "activity_edit"]
    assert {str(a.target_id) for a in denied} == {act["id"]}

    # 작성자 수정 — 새로 넣는 참석자만 소속 검사, deal_id "" 는 연결 해제
    r = await client.patch(url, json={"contact_ids": [c1["id"], outsider["id"]]}, headers=h)
    assert r.status_code == 422
    r = await client.patch(url, json={"deal_id": other_deal["id"]}, headers=h)
    assert r.status_code == 422
    r = await client.patch(
        url, json={"occurred_at": (datetime.now(UTC) + timedelta(hours=1)).isoformat()}, headers=h
    )
    assert r.status_code == 422
    assert (await client.patch(url, json={"body": None}, headers=h)).status_code == 422
    r = await client.patch(
        url,
        json={"deal_id": "", "contact_ids": [c1["id"]], "user_ids": [], "kind": "email"},
        headers=h,
    )
    assert r.status_code == 200, r.text
    edited = r.json()
    assert (edited["deal"], edited["kind"], edited["users"]) == (None, "email", [])
    assert [p["id"] for p in edited["contacts"]] == [c1["id"]]
    assert edited["body"] == "2호 펀드 제안서 설명."  # 안 보낸 칸은 그대로
    # 이직한 참석자는 그대로 둘 수 있다
    await client.patch(f"{API}/contacts/{c1['id']}", json={"org_id": other["id"]}, headers=h)
    r = await client.patch(url, json={"contact_ids": [c1["id"], c2["id"]]}, headers=h)
    assert r.status_code == 200, r.text

    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        r = await boss.patch(url, json={"body": "임원이 고침"}, headers=hx)
        assert (r.status_code, r.json()["can_edit"], r.json()["author"]["name"]) == (
            200,
            True,
            "김수진",
        )
    assert (await client.delete(url, headers=h)).status_code == 204
    assert (await client.delete(url, headers=h)).status_code == 404
    assert (await client.patch(url, json={"body": "x"}, headers=h)).status_code == 404
    by_org = (await client.get(f"{API}/activities", params={"org_id": org["id"]})).json()["items"]
    assert [a["id"] for a in by_org] == [plain["id"]]
    [ev] = await _audits(app_db, "crm.activity_deleted")
    assert (str(ev.target_id), ev.actor_id, ev.detail) == (act["id"], users.member.id, None)


CSV_HEADER = "기관,유형,이름,부서·직함,전화,이메일,메모\r\n"


async def test_import(client: AsyncClient, users: Users, app_db) -> None:
    """dry_run 미리보기 → 반영. 행별 ok·skip·error, 정규화 이름으로 기존 기관에 붙임, 감사는 반영 때만."""
    h = await _as(client, "sujin@zenith.co.kr")
    existing = await _org(client, h, "(주)미래인프라자산운용", "am")
    await _contact(client, h, existing["id"], "한도윤")
    rows = (
        CSV_HEADER
        + "미래인프라 자산운용,,한도윤,상무,,,\r\n"  # 2: 기존 기관·기존 담당자 → skip
        + "미래인프라 자산운용,,김새봄,팀장,02-1,sb@mirae.example,\r\n"  # 3: 기존 기관에 붙임
        + ",,,,,,\r\n"  # 4: 빈 행은 버리되 줄 번호는 센다
        + '세움은행,은행,"이, 은행",과장,,,"여러 줄\n메모"\r\n'  # 5: 새 기관 + 담당자(따옴표·쉼표·줄바꿈)
        + "세움은행,은행,이 은행,,,,\r\n"  # 6: 같은 파일에서 앞서 받은 담당자 아님(이름 다름) → ok
        + '세움 은행,,"이, 은행",,,,\r\n'  # 7: 같은 파일에서 받은 담당자 → skip
        + "누리공제회,공제회,,,,,연간 출자 계획\r\n"  # 8: 기관만 → 새 기관(메모)
        + "누리공제회,,,,,,\r\n"  # 9: 이미 있는 기관
        + "다온,우주,,,,,\r\n"  # 10: 모르는 유형
        + "다온,,,차장,,,\r\n"  # 11: 이름 없이 직함
        + "다온,,정유나,,,not-an-email,\r\n"  # 12: 이메일 형식
        + ",,홍길동,,,,\r\n"  # 13: 기관 비었음
    )
    data = ("﻿" + rows).encode("utf-8")

    async def send(dry: str, payload: bytes = data, name: str = "a.csv"):
        return await client.post(
            f"{API}/orgs/import",
            params={"dry_run": dry},
            files={"file": (name, payload, "text/csv")},
            headers=h,
        )

    r = await send("true")
    assert r.status_code == 200, r.text
    preview = r.json()
    got = [(row["line"], row["status"], row["note"]) for row in preview["rows"]]
    assert got == [
        (2, "skip", "이미 있는 담당자"),
        (3, "ok", "기존 기관 「(주)미래인프라자산운용」 에 붙임"),
        (5, "ok", "새 기관"),
        (6, "ok", ""),
        (7, "skip", "이미 있는 담당자"),
        (8, "ok", "새 기관"),
        (9, "skip", "이미 있는 기관"),
        (10, "error", "유형 「우주」 을 알 수 없습니다."),
        (11, "error", "담당자 이름이 비었습니다."),
        (12, "error", "이메일 형식이 아닙니다."),
        (13, "error", "기관 이름은 1자 이상 100자 이하여야 합니다."),
    ]
    assert (preview["new_orgs"], preview["new_contacts"]) == (2, 3)
    assert preview["rows"][2]["contact_name"] == "이, 은행"
    assert (
        len((await client.get(f"{API}/orgs")).json()["items"]) == 1
    )  # 미리보기는 아무것도 안 만든다
    assert await _audits(app_db, "crm.imported") == []

    r = await send("false")
    assert r.json() == preview  # 반영도 같은 모양
    orgs = {o["name"]: o for o in (await client.get(f"{API}/orgs")).json()["items"]}
    assert set(orgs) == {"(주)미래인프라자산운용", "세움은행", "누리공제회"}
    assert (orgs["세움은행"]["org_type"], orgs["세움은행"]["contact_count"]) == ("bank", 2)
    nuri = (await client.get(f"{API}/orgs/{orgs['누리공제회']['id']}")).json()
    assert (nuri["memo"], nuri["contacts"], nuri["owners"]) == ("연간 출자 계획", [], [])
    seum = (await client.get(f"{API}/orgs/{orgs['세움은행']['id']}")).json()
    assert {c["name"]: c["memo"] for c in seum["contacts"]} == {
        "이, 은행": "여러 줄\n메모",
        "이 은행": "",
    }
    assert orgs["(주)미래인프라자산운용"]["contact_count"] == 2
    [ev] = await _audits(app_db, "crm.imported")
    assert (ev.target_id, ev.detail) == (None, {"orgs": 2, "contacts": 3})
    # 다시 올리면 전부 skip·error
    again = (await send("false")).json()
    assert (again["new_orgs"], again["new_contacts"]) == (0, 0)

    # CP949(한글 윈도우 엑셀) 도 읽는다
    r = await send("true", (CSV_HEADER + "한결연금공단,연기금,박서준,,,,\r\n").encode("cp949"))
    assert r.json()["rows"][0]["status"] == "ok"
    # 파일 단위 거부 — 전부 422
    for payload in (
        b"",
        "이름,전화\r\n박서준,1\r\n".encode(),  # 「기관」 칸 없음
        CSV_HEADER.encode(),  # 본문 0행
        (CSV_HEADER + '한결,,"닫히지 않음\r\n').encode(),
        (CSV_HEADER + "한결,,,,,,\r\n" * 2001).encode(),
        b"x" * (1024 * 1024 + 1),
    ):
        r = await send("true", payload)
        assert (r.status_code, r.json()["code"]) == (422, "VALIDATION"), payload[:40]
    r = await client.post(
        f"{API}/orgs/import", params={"dry_run": "true"}, headers=h, data={"x": "1"}
    )
    assert (r.status_code, r.json()["detail"]) == (422, "CSV 파일을 올려 주세요.")


async def test_export(client: AsyncClient, users: Users, app_db) -> None:
    """xlsx 로 열리고 머리글·정렬·제외(erased·지운 기관)·수식 아님. 감사 crm.exported {entity, rows}."""
    h = await _as(client, "sujin@zenith.co.kr")
    org = await _org(
        client,
        h,
        "한결연금공단",
        memo='=HYPERLINK("http://x")',
        owner_ids=[str(users.lead.id), str(users.member.id)],
    )
    other = await _org(client, h, "누리공제회", "mutual_aid")
    await _contact(client, h, org["id"], "박서준", title="팀장", email="sj@hangyeol.example")
    await _contact(client, h, other["id"], "최민호")
    await _activity(client, h, org["id"])
    await _deal(client, h, org["id"], amount_krw=50_000_000_000, owner_id=str(users.lead.id))
    await _deal(client, h, other["id"], kind="disposal", name="건물 매각")

    async def sheet(entity: str):
        r = await client.post(f"{API}/export", params={"entity": entity}, headers=h)
        assert r.status_code == 200, r.text
        assert (
            r.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "content-disposition" not in r.headers
        ws = load_workbook(io.BytesIO(r.content)).active
        return ws, [list(row) for row in ws.iter_rows(values_only=True)]

    ws, rows = await sheet("orgs")
    assert ws.title == "기관"
    assert ws.freeze_panes == "B2"
    assert ws["A1"].font.bold
    assert rows[0] == ["기관", "유형", "사내 창구", "담당자 수", "마지막 연락", "메모", "등록일"]
    assert [r[0] for r in rows[1:]] == ["누리공제회", "한결연금공단"]
    hangyeol = rows[2]
    assert (hangyeol[1], hangyeol[2], hangyeol[3]) == ("연기금", "김수진, 박동현", 1)
    assert hangyeol[5] == '=HYPERLINK("http://x")'
    assert ws["F3"].data_type == "s"  # 수식이 아니라 글자
    assert len(hangyeol[4]) == len("2026-10-02")

    _, rows = await sheet("contacts")
    assert rows == [
        ["기관", "이름", "부서·직함", "전화", "이메일", "메모"],
        ["누리공제회", "최민호", None, None, None, None],
        ["한결연금공단", "박서준", "팀장", None, "sj@hangyeol.example", None],
    ]

    ws, rows = await sheet("deals")
    assert rows[0][5] == "금액(원)"
    assert [r[0] for r in rows[1:]] == ["매각", "출자 유치"]  # 등록 최신순
    assert rows[2][1:7] == [
        "인프라 2호 펀드 출자",
        "한결연금공단",
        "접촉",
        "진행 중",
        50_000_000_000,
        "박동현",
    ]
    assert ws["F3"].number_format == "#,##0"

    events = await _audits(app_db, "crm.exported")
    assert [e.detail for e in events] == [
        {"entity": "orgs", "rows": 2},
        {"entity": "contacts", "rows": 2},
        {"entity": "deals", "rows": 2},
    ]
    r = await client.post(f"{API}/export", params={"entity": "users"}, headers=h)
    assert (r.status_code, r.json()["code"]) == (422, "VALIDATION")
    assert len(await _audits(app_db, "crm.exported")) == 3


async def test_review_followups(
    client: AsyncClient, make_client: Callable[[], AsyncClient], users: Users, app_db
) -> None:
    """리뷰 후속: lead 도 삭제·erase 403, 상한 422, erased 담당자를 새 활동에 못 넣음, 지운 기관의 안건·활동·담당자."""
    h = await _as(client, "sujin@zenith.co.kr")
    org = await _org(client, h, "한결연금공단")
    keep = await _org(client, h, "누리공제회", "mutual_aid")
    c1 = await _contact(client, h, org["id"], "박서준")
    c2 = await _contact(client, h, org["id"], "이하은")
    act = await _activity(client, h, org["id"], contact_ids=[c1["id"]])

    async with make_client() as lead:
        hl = await _as(lead, "donghyun@zenith.co.kr")
        r = await lead.request(
            "DELETE", f"{API}/orgs/{org['id']}", json={"reason": "x"}, headers=hl
        )
        assert r.status_code == 403
        assert (await lead.post(f"{API}/contacts/{c1['id']}/erase", headers=hl)).status_code == 403
    denied = await _audits(app_db, "acl.denied")
    assert [(a.actor_id, a.detail["reason"]) for a in denied] == [
        (users.lead.id, "org_delete"),
        (users.lead.id, "contact_erase"),
    ]

    # 상한
    assert (await client.get(f"{API}/orgs", params={"stale_days": 36501})).status_code == 422
    many = [str(users.member.id)] * 201
    r = await client.post(
        f"{API}/orgs", json={"name": "x", "org_type": "etc", "owner_ids": many}, headers=h
    )
    assert r.status_code == 422
    r = await client.patch(f"{API}/activities/{act['id']}", json={"user_ids": many}, headers=h)
    assert r.status_code == 422
    for amount in ("1000", 1.0, True):
        body = {"kind": "investment", "name": "x", "org_id": org["id"], "amount_krw": amount}
        assert (await client.post(f"{API}/deals", json=body, headers=h)).status_code == 422, amount

    # erased 담당자는 새 활동에 못 넣는다 — 이미 참석자였던 기록은 그대로 고칠 수 있다
    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        assert (await boss.post(f"{API}/contacts/{c1['id']}/erase", headers=hx)).status_code == 200
    body = {"kind": "call", "occurred_at": _ago(0), "org_id": org["id"], "body": "x"}
    r = await client.post(f"{API}/activities", json={**body, "contact_ids": [c1["id"]]}, headers=h)
    assert r.status_code == 422
    url = f"{API}/activities/{act['id']}"
    r = await client.patch(url, json={"contact_ids": [c1["id"], c2["id"]]}, headers=h)
    assert r.status_code == 200, r.text
    assert [p["name"] for p in r.json()["contacts"]] == ["", "이하은"]

    # CSV 의 NUL 은 버린다
    csv_body = "기관,이름\r\n세움\x00은행,김\x00새봄\r\n".encode()
    r = await client.post(
        f"{API}/orgs/import",
        params={"dry_run": "true"},
        files={"file": ("a.csv", csv_body, "text/csv")},
        headers=h,
    )
    assert r.json()["rows"][0] | {"note": ""} == {
        "line": 2,
        "org_name": "세움은행",
        "contact_name": "김새봄",
        "status": "ok",
        "note": "",
    }

    # 지운 기관: 안건은 deals 내보내기에서, 활동·담당자 수정은 404
    gone = await _deal(client, h, org["id"], name="지울 안건")
    await _deal(client, h, keep["id"], name="남는 안건")
    await client.post(
        f"{API}/deals/{gone['id']}/drop", json={"reason": "x", "version": 1}, headers=h
    )
    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        r = await boss.request(
            "DELETE", f"{API}/orgs/{org['id']}", json={"reason": "정리"}, headers=hx
        )
        assert r.status_code == 204
        r = await boss.post(f"{API}/export", params={"entity": "deals"}, headers=hx)
        rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
        assert [row[1] for row in rows[1:]] == ["남는 안건"]
        assert (await boss.patch(url, json={"body": "x"}, headers=hx)).status_code == 404
        assert (await boss.delete(url, headers=hx)).status_code == 404
    r = await client.patch(f"{API}/contacts/{c2['id']}", json={"memo": "x"}, headers=h)
    assert r.status_code == 404


async def test_contact_and_deal_delete(
    client: AsyncClient, make_client: Callable[[], AsyncClient], users: Users, app_db
) -> None:
    """담당자·안건 지우기(목록에서 뺌): 등록한 사람은 활동이 붙기 전까지, exec 는 언제나. 403 + 감사, 지운 뒤 가시성."""
    h = await _as(client, "sujin@zenith.co.kr")
    org = await _org(client, h, "한결연금공단")
    c = await _contact(client, h, org["id"], "박서준")
    spare = await _contact(client, h, org["id"], "이하은")
    assert c["can_delete"] is True
    deal = await _deal(client, h, org["id"])
    assert deal["can_delete"] is True
    quick = await _deal(client, h, org["id"], name="잘못 만든 안건")

    lead = make_client()
    hl = await _as(lead, "donghyun@zenith.co.kr")
    seen = (await lead.get(f"{API}/orgs/{org['id']}")).json()["contacts"]
    assert [x["can_delete"] for x in seen] == [False, False]
    assert (await lead.get(f"{API}/deals/{deal['id']}")).json()["can_delete"] is False
    assert (await lead.delete(f"{API}/contacts/{c['id']}", headers=hl)).status_code == 403
    assert (await lead.delete(f"{API}/deals/{deal['id']}", headers=hl)).status_code == 403
    denied = await _audits(app_db, "acl.denied")
    assert sorted(a.detail["reason"] for a in denied) == ["contact_delete", "deal_delete"]

    # 등록한 사람 — 활동이 붙기 전
    assert (await client.delete(f"{API}/deals/{quick['id']}", headers=h)).status_code == 204
    [ev] = await _audits(app_db, "crm.deal_deleted")
    assert (str(ev.target_id), ev.detail) == (quick["id"], {"name": "잘못 만든 안건"})
    # 활동이 붙으면 등록한 사람도 못 지운다
    act = await _activity(client, h, org["id"], deal_id=deal["id"], contact_ids=[c["id"]])
    detail = (await client.get(f"{API}/orgs/{org['id']}")).json()
    assert {x["name"]: x["can_delete"] for x in detail["contacts"]} == {
        "박서준": False,
        "이하은": True,
    }
    assert (await client.get(f"{API}/deals/{deal['id']}")).json()["can_delete"] is False
    assert (await client.delete(f"{API}/contacts/{c['id']}", headers=h)).status_code == 403
    assert (await client.delete(f"{API}/deals/{deal['id']}", headers=h)).status_code == 403
    # 활동에 안 붙은 담당자는 등록한 사람이 지운다
    assert (await client.delete(f"{API}/contacts/{spare['id']}", headers=h)).status_code == 204

    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        assert (await boss.delete(f"{API}/contacts/{c['id']}", headers=hx)).status_code == 204
        assert (await boss.delete(f"{API}/contacts/{c['id']}", headers=hx)).status_code == 404
        assert (await boss.delete(f"{API}/deals/{deal['id']}", headers=hx)).status_code == 204
        assert (await boss.delete(f"{API}/deals/{deal['id']}", headers=hx)).status_code == 404
        # 목록에서 지운 담당자도 개인정보 삭제는 된다
        assert (await boss.post(f"{API}/contacts/{c['id']}/erase", headers=hx)).status_code == 200
        contacts_xlsx = await boss.post(f"{API}/export", params={"entity": "contacts"}, headers=hx)
        assert load_workbook(io.BytesIO(contacts_xlsx.content)).active.max_row == 1
    await lead.aclose()
    events = await _audits(app_db, "crm.contact_deleted")
    assert sorted(str(e.target_id) for e in events) == sorted([spare["id"], c["id"]])
    assert all(e.detail is None for e in events)

    # 담당자: 기관 상세·수·수정·새 활동에서 빠지고, 옛 활동 참석자로는 남는다
    detail = (await client.get(f"{API}/orgs/{org['id']}")).json()
    assert (detail["contacts"], detail["contact_count"]) == ([], 0)
    assert (
        await client.patch(f"{API}/contacts/{spare['id']}", json={"memo": "x"}, headers=h)
    ).status_code == 404
    body = {"kind": "call", "occurred_at": _ago(0), "org_id": org["id"], "body": "x"}
    r = await client.post(
        f"{API}/activities", json={**body, "contact_ids": [spare["id"]]}, headers=h
    )
    assert r.status_code == 422
    timeline = (await client.get(f"{API}/activities", params={"org_id": org["id"]})).json()["items"]
    assert [x["id"] for x in timeline[0]["contacts"]] == [c["id"]]
    # 안건: 목록·상세·수정·타임라인에서 빠지고, 붙어 있던 활동은 기관 타임라인에 안건 표시 없이 남는다
    assert timeline[0]["id"] == act["id"] and timeline[0]["deal"] is None
    assert (await client.get(f"{API}/deals")).json()["items"] == []
    assert (await client.get(f"{API}/deals/{deal['id']}")).status_code == 404
    r = await client.patch(f"{API}/deals/{deal['id']}", json={"version": 1, "name": "x"}, headers=h)
    assert r.status_code == 404
    assert (
        await client.get(f"{API}/activities", params={"deal_id": deal["id"]})
    ).status_code == 404
    r = await client.post(f"{API}/activities", json={**body, "deal_id": deal["id"]}, headers=h)
    assert r.status_code == 422
    # 지운 안건에 걸려 있던 활동을 안건 칸 그대로(낡은 화면) 고쳐도 막히지 않는다
    r = await client.patch(f"{API}/activities/{act['id']}", json={"body": "고침"}, headers=h)
    assert r.status_code == 200
    # 지운 진행 중 안건은 기관 삭제를 막지 않는다
    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        r = await boss.request(
            "DELETE", f"{API}/orgs/{org['id']}", json={"reason": "정리"}, headers=hx
        )
        assert r.status_code == 204


async def test_deleted_orgs_and_restore(
    client: AsyncClient, make_client: Callable[[], AsyncClient], users: Users, app_db
) -> None:
    """지운 기관 목록(사유·같은 이름 기관)과 되살리기 — exec 만(403 + 감사), 담당자·안건이 그대로 돌아온다."""
    h = await _as(client, "sujin@zenith.co.kr")
    org = await _org(client, h, "(주)미래인프라")
    await _contact(client, h, org["id"], "한도윤")
    deal = await _deal(client, h, org["id"])
    await client.post(
        f"{API}/deals/{deal['id']}/drop", json={"reason": "x", "version": 1}, headers=h
    )
    url = f"{API}/orgs/{org['id']}"

    assert (await client.get(f"{API}/orgs/deleted")).status_code == 403
    assert (await client.post(f"{url}/restore", headers=h)).status_code == 403
    denied = await _audits(app_db, "acl.denied")
    assert [a.detail["reason"] for a in denied] == ["org_deleted_list", "org_restore"]

    async with make_client() as boss:
        hx = await _as(boss, "jaeho@zenith.co.kr")
        assert (await boss.get(f"{API}/orgs/deleted")).json()["items"] == []
        assert (await boss.post(f"{url}/restore", headers=hx)).status_code == 404  # 안 지운 기관
        r = await boss.request("DELETE", url, json={"reason": " 중복 등록 "}, headers=hx)
        assert r.status_code == 204
        twin = await _org(boss, hx, "미래 인프라")
        [gone] = (await boss.get(f"{API}/orgs/deleted")).json()["items"]
        assert (gone["id"], gone["name"], gone["org_type"], gone["reason"]) == (
            org["id"],
            "(주)미래인프라",
            "pension",
            "중복 등록",
        )
        assert gone["duplicate"] == {"id": twin["id"], "name": "미래 인프라"}
        assert gone["deleted_at"]

        r = await boss.post(f"{url}/restore", headers=hx)
        assert r.status_code == 200, r.text
        back = r.json()
        assert ([c["name"] for c in back["contacts"]], back["contact_count"]) == (["한도윤"], 1)
        assert (await boss.post(f"{url}/restore", headers=hx)).status_code == 404  # 두 번째는 없다
        assert (await boss.get(f"{API}/orgs/deleted")).json()["items"] == []
        missing = f"{API}/orgs/00000000-0000-0000-0000-000000000000/restore"
        assert (await boss.post(missing, headers=hx)).status_code == 404
    [ev] = await _audits(app_db, "crm.org_restored")
    assert (str(ev.target_id), ev.detail) == (org["id"], {"name": "(주)미래인프라"})
    assert (await client.get(url)).status_code == 200
    assert [d["id"] for d in (await client.get(f"{API}/deals")).json()["items"]] == [deal["id"]]
    names = [o["name"] for o in (await client.get(f"{API}/orgs")).json()["items"]]
    assert names == ["(주)미래인프라", "미래 인프라"]


async def test_import_undo(
    client: AsyncClient, make_client: Callable[[], AsyncClient], users: Users, app_db
) -> None:
    """가져오기 되돌리기: 가져온 사람만(403 + 감사)·24시간 안(409). 쓰인 기관·담당자는 남기고, 두 번째는 404."""
    h = await _as(client, "sujin@zenith.co.kr")
    existing = await _org(client, h, "기존기관")
    assert (await client.get(f"{API}/imports/last")).json() == {"import": None}
    csv_body = (
        CSV_HEADER
        + "가기관,,가1,,,,\r\n"
        + "가기관,,가2,,,,\r\n"
        + "나기관,,나1,,,,\r\n"
        + "다기관,,,,,,메모만\r\n"
        + "기존기관,,기1,,,,\r\n"
    ).encode()

    async def send(dry: str):
        return await client.post(
            f"{API}/orgs/import",
            params={"dry_run": dry},
            files={"file": ("a.csv", csv_body, "text/csv")},
            headers=h,
        )

    await send("true")
    assert (await client.get(f"{API}/imports/last")).json() == {
        "import": None
    }  # 미리보기는 안 남는다
    await send("false")
    last = (await client.get(f"{API}/imports/last")).json()["import"]
    assert (last["new_orgs"], last["new_contacts"]) == (3, 4)

    orgs = {o["name"]: o for o in (await client.get(f"{API}/orgs")).json()["items"]}
    ga = (await client.get(f"{API}/orgs/{orgs['가기관']['id']}")).json()
    na = (await client.get(f"{API}/orgs/{orgs['나기관']['id']}")).json()
    assert ga["contacts"][0]["can_delete"] is True  # 가져온 사람이 등록한 사람이다
    # 가기관은 쓰이기 시작했다(활동), 나기관의 담당자는 목록에서 지웠다
    await _activity(client, h, ga["id"], contact_ids=[ga["contacts"][0]["id"]])
    assert (
        await client.delete(f"{API}/contacts/{na['contacts'][0]['id']}", headers=h)
    ).status_code == 204

    undo = f"{API}/imports/{last['id']}/undo"
    async with make_client() as lead:
        hl = await _as(lead, "donghyun@zenith.co.kr")
        assert (await lead.get(f"{API}/imports/last")).json() == {"import": None}  # 내 것만
        assert (await lead.post(undo, headers=hl)).status_code == 403
    [denied] = await _audits(app_db, "acl.denied")
    assert (denied.detail["reason"], str(denied.target_id)) == ("import_undo", last["id"])

    r = await client.post(undo, headers=h)
    assert r.status_code == 200, r.text
    # 지운 것: 나기관·다기관, 기1 (나1 은 이미 목록에서 지운 행이라 수에 안 넣는다). 남긴 것: 가기관과 가1·가2
    assert r.json() == {
        "removed_orgs": 2,
        "removed_contacts": 1,
        "kept_orgs": 1,
        "kept_contacts": 2,
    }
    [ev] = await _audits(app_db, "crm.import_undone")
    assert (ev.target_id, ev.detail) == (None, r.json())
    names = [o["name"] for o in (await client.get(f"{API}/orgs")).json()["items"]]
    assert names == ["가기관", "기존기관"]
    assert (await client.get(f"{API}/orgs/{existing['id']}")).json()["contacts"] == []
    assert [c["name"] for c in (await client.get(f"{API}/orgs/{ga['id']}")).json()["contacts"]] == [
        "가1",
        "가2",
    ]

    # 멱등 — 두 번째는 404, 더 지우지 않는다
    assert (await client.post(undo, headers=h)).status_code == 404
    assert len((await client.get(f"{API}/orgs")).json()["items"]) == 2
    assert (await client.get(f"{API}/imports/last")).json() == {"import": None}
    missing = f"{API}/imports/00000000-0000-0000-0000-000000000000/undo"
    assert (await client.post(missing, headers=h)).status_code == 404

    # 24시간이 지나면 last 에서 빠지고 되돌리기는 409
    old = CrmImport(user_id=users.member.id, created_at=datetime.now(UTC) - timedelta(hours=25))
    app_db.add(old)
    await app_db.flush()
    assert (await client.get(f"{API}/imports/last")).json() == {"import": None}
    r = await client.post(f"{API}/imports/{old.id}/undo", headers=h)
    assert (r.status_code, r.json()["code"]) == (409, "CONFLICT")
