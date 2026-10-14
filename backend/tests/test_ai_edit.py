"""AI 파일 수정 (README 53) — 편집 연산 적용기와 `POST /ai/attachments/{id}/edit`."""

import io
import json
import zipfile
from collections.abc import AsyncIterator

import openpyxl
import pytest
from fastapi import FastAPI
from httpx import AsyncClient
from sqlalchemy import select

from app.ai import edit
from app.ai.provider import get_provider
from app.models import AiAttachment, AuditEvent
from app.routers import ai

from .conftest import Users, csrf, login
from .test_ai import FakeProvider
from .test_ai_attachments import _upload, docx_bytes, xlsx_bytes

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def test_xlsx_dump_and_apply_preserve_the_rest() -> None:
    data = xlsx_bytes()
    assert (
        edit.dump_xlsx(data)
        == "## 매출\n매출!A1 = 월\n매출!B1 = 매출\n매출!A2 = 1월\n매출!B2 = 120\n## 메모\n메모!A1 = 비고"
    )
    out, changes = edit.apply_xlsx(
        data,
        [
            {"sheet": "매출", "cell": "b2", "value": 1250},
            {"sheet": "매출", "cell": "C2", "value": "=B2*2"},
        ],
    )
    assert [(c.where, c.before, c.after) for c in changes] == [
        ("매출!B2", "120", "1250"),
        ("매출!C2", "", "=B2*2"),
    ]
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert (
        wb["매출"]["B2"].value == 1250
        and wb["매출"]["C2"].value == "=B2*2"
        and wb["매출"]["A1"].value == "월"
    )
    assert wb["메모"]["A1"].value == "비고"  # 다른 시트 그대로
    for bad in (
        [{"sheet": "없음", "cell": "A1", "value": 1}],
        [{"sheet": "매출", "cell": "A", "value": 1}],
        [{"sheet": "매출", "cell": "A1", "value": [1]}],
    ):
        with pytest.raises(edit.BadOp):
            edit.apply_xlsx(data, bad)


def test_docx_dump_and_apply_keep_formatting_and_prefixes() -> None:
    xml = (
        f'<?xml version="1.0"?><w:document xmlns:w="{W}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" mc:Ignorable="w14"><w:body>'
        '<w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:b/></w:rPr><w:t>제목</w:t></w:r><w:r><w:t> 꼬리</w:t></w:r></w:p>'
        "<w:p><w:r><w:t>본문</w:t></w:r></w:p>"
        "<w:p/></w:body></w:document>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
        z.writestr("word/styles.xml", "<styles/>")
    data = buf.getvalue()
    assert edit.dump_docx(data) == "[p0] 제목 꼬리\n[p1] 본문\n[p2]"
    out, changes = edit.apply_docx(
        data, [{"paragraph": 0, "text": "새 제목\t탭\n둘째 줄"}, {"paragraph": 2, "text": "채움"}]
    )
    assert [(c.where, c.before, c.after) for c in changes] == [
        ("문단 0", "제목 꼬리", "새 제목\t탭\n둘째 줄"),
        ("문단 2", "", "채움"),
    ]
    doc = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode()
    assert 'xmlns:w="' in doc and "ns0:" not in doc and 'mc:Ignorable="w14"' in doc  # 접두사 보존
    assert (
        "<w:b/>" in doc or "<w:b />" in doc
    ) and 'w:val="center"' in doc  # 첫 run·문단 서식 유지
    assert edit.dump_docx(out) == "[p0] 새 제목 탭 둘째 줄\n[p1] 본문\n[p2] 채움"
    assert (
        zipfile.ZipFile(io.BytesIO(out)).read("word/styles.xml") == b"<styles/>"
    )  # 다른 항목 그대로
    with pytest.raises(edit.BadOp):
        edit.apply_docx(data, [{"paragraph": 9, "text": "x"}])


class Editing(FakeProvider):
    reply = "[]"

    async def stream(
        self, *, system, messages, max_tokens=None, tools=None, tool_choice=None
    ) -> AsyncIterator:
        Editing.seen = messages
        Editing.system = system
        yield self.reply


@pytest.mark.usefixtures("minio_bucket")
async def test_edit_endpoint_makes_new_attachment_with_changes(
    app: FastAPI, client: AsyncClient, users: Users, app_db, monkeypatch
) -> None:
    Editing.reply = json.dumps([{"sheet": "매출", "cell": "B2", "value": 999}])
    app.dependency_overrides[get_provider] = Editing
    ai._upload_calls.clear()
    ai._check_calls.clear()
    await login(client, "sujin@zenith.co.kr")
    att = (await _upload(client, "매출.xlsx", xlsx_bytes())).json()
    r = await client.post(
        f"/api/ai/attachments/{att['id']}/edit",
        json={"instruction": "1월 매출을 999로"},
        headers=csrf(client),
    )
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["changes"] == [{"where": "매출!B2", "before": "120", "after": "999"}]
    assert out["result"]["name"] == "매출 (AI 수정).xlsx" and out["result"]["id"] != att["id"]
    assert (
        Editing.system[1]["text"] == ai.EDIT_PROMPT
        and "매출!B2 = 120" in Editing.seen[0]["content"]
    )
    assert "지시: 1월 매출을 999로" in Editing.seen[0]["content"]
    d = await client.get(f"/api/ai/attachments/{out['result']['id']}")
    assert (
        d.status_code == 200
        and openpyxl.load_workbook(io.BytesIO(d.content))["매출"]["B2"].value == 999
    )
    orig = await client.get(f"/api/ai/attachments/{att['id']}")
    assert (
        openpyxl.load_workbook(io.BytesIO(orig.content))["매출"]["B2"].value == 120
    )  # 원본 그대로
    assert (
        await app_db.scalar(select(AuditEvent).where(AuditEvent.kind == "ai.attachment_edit"))
        is not None
    )
    saved_text = await app_db.scalar(
        select(AiAttachment.text).where(AiAttachment.id == out["result"]["id"])
    )
    assert saved_text.startswith("## 매출")
    # docx 도
    Editing.reply = json.dumps([{"paragraph": 0, "text": "바뀐 첫 문단"}])
    att2 = (await _upload(client, "글.docx", docx_bytes(["첫 문단", "둘째"]))).json()
    r = await client.post(
        f"/api/ai/attachments/{att2['id']}/edit",
        json={"instruction": "첫 문단 고쳐"},
        headers=csrf(client),
    )
    assert r.status_code == 200 and r.json()["changes"][0]["after"] == "바뀐 첫 문단"
    ai._check_calls.clear()


@pytest.mark.usefixtures("minio_bucket")
async def test_edit_rejections(
    app: FastAPI, client: AsyncClient, users: Users, app_db, monkeypatch
) -> None:
    app.dependency_overrides[get_provider] = Editing
    ai._upload_calls.clear()
    ai._check_calls.clear()
    await login(client, "sujin@zenith.co.kr")
    txt = (await _upload(client, "메모.txt", b"hello")).json()
    r = await client.post(
        f"/api/ai/attachments/{txt['id']}/edit", json={"instruction": "x"}, headers=csrf(client)
    )
    assert r.status_code == 422 and r.json()["detail"] == "xlsx·docx 만 고칠 수 있습니다."
    xl = (await _upload(client, "매출.xlsx", xlsx_bytes())).json()
    Editing.reply = "표로 정리하면"  # 배열 아님
    r = await client.post(
        f"/api/ai/attachments/{xl['id']}/edit", json={"instruction": "x"}, headers=csrf(client)
    )
    assert r.status_code == 503 and r.json()["code"] == "AI_FORMAT"
    Editing.reply = json.dumps([{"sheet": "없는시트", "cell": "A1", "value": 1}])
    r = await client.post(
        f"/api/ai/attachments/{xl['id']}/edit", json={"instruction": "x"}, headers=csrf(client)
    )
    assert r.status_code == 503 and "없는 시트" in r.json()["detail"]
    await login(client, "jaeho@zenith.co.kr")
    r = await client.post(
        f"/api/ai/attachments/{xl['id']}/edit", json={"instruction": "x"}, headers=csrf(client)
    )
    assert r.status_code == 404  # 남의 첨부
    assert (
        await app_db.scalar(select(AiAttachment).where(AiAttachment.name.like("%(AI 수정)%")))
        is None
    )
    ai._check_calls.clear()
