"""AI 파일 만들기(README 55) — 모델이 낸 표(시트·행)를 xlsx 로 쓴다. 입력은 신뢰하지 않는다(모델 출력 = 사용자 글·검색
근거에서 파생) — 크기 상한을 넘거나 모양이 틀리면 BadSpec. 서식은 머리글 굵게 하나뿐(ponytail: 열 너비·숫자 서식은 요청이
오면). 수식은 허용 목록 함수만, 문자열·외부 참조 없이 — `=WEBSERVICE("…"&B2)`·`=HYPERLINK` 로 셀 값을 밖으로 내보내거나
DDE 를 심지 못하게(보안 리뷰, 루트 README 가드레일)."""

import io
import math
import re
from typing import Any

import openpyxl
from openpyxl.styles import Font

MAX_SHEETS = 10
MAX_ROWS = 2000
MAX_COLS = 50
MAX_CELLS = 20_000  # 전체 셀 수 — 2000행 × 50열을 다 채우지는 못하게
MAX_CELL_CHARS = 1000
MAX_TITLE_CHARS = 31  # 엑셀 시트 이름 상한
_BAD_TITLE = set("[]:*?/\\")  # 엑셀이 시트 이름에 허락하지 않는 글자
_RESERVED_TITLES = {"history"}  # 엑셀 예약 시트 이름
# 수식 허용 목록 — 계산 함수만. 문자열 리터럴(")·외부 통합 문서([)·DDE(|)·다른 시트 참조(')는 통째로 거절한다
_SAFE_FUNCS = {
    "SUM",
    "AVERAGE",
    "MIN",
    "MAX",
    "COUNT",
    "COUNTA",
    "ROUND",
    "ROUNDUP",
    "ROUNDDOWN",
    "IF",
    "ABS",
    "SUMIF",
    "COUNTIF",
    "AVERAGEIF",
    "SUMPRODUCT",
    "INT",
    "MOD",
    "AND",
    "OR",
    "NOT",
    "IFERROR",
}
_FUNC_RE = re.compile(r"([A-Z][A-Z0-9.]*)\s*\(", re.I)
_FORMULA_BAD_CHARS = set("\"'|[]!")


class BadSpec(ValueError):
    """모델이 낸 표가 모양·상한을 어겼다 — 도구 결과로 돌려주면 모델이 고쳐 다시 부른다."""


def _xml_ok(ch: str) -> bool:
    # xlsx(XML 1.0)가 받는 글자만 — 제어문자·고립 서로게이트·U+FFFE/FFFF 는 openpyxl 이 저장하다 터지거나 깨진 파일이 된다
    return ch in "\n\t" or " " <= ch <= "\ud7ff" or "\ue000" <= ch <= "\ufffd" or ch > "\uffff"


def build_xlsx(sheets: Any) -> bytes:
    """`[{"title": 시트, "rows": [[값…]…]}…]` → xlsx 바이트. 값은 문자열·숫자·불리언·null, "=" 로 시작하면 수식."""
    if not isinstance(sheets, list) or not sheets:
        raise BadSpec("sheets 는 비어 있지 않은 배열이어야 합니다.")
    if len(sheets) > MAX_SHEETS:
        raise BadSpec(f"시트는 {MAX_SHEETS}개까지입니다.")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cells = 0
    seen: set[str] = set()
    for i, sheet in enumerate(sheets, start=1):
        if not isinstance(sheet, dict):
            raise BadSpec("시트는 {title, rows} 객체여야 합니다.")
        title = sheet.get("title")
        if not isinstance(title, str) or not title.strip():
            title = f"Sheet{i}"
        title = "".join("_" if c in _BAD_TITLE or not _xml_ok(c) else c for c in title)
        title = title.strip().strip("'")[:MAX_TITLE_CHARS].strip() or f"Sheet{i}"
        if title.lower() in _RESERVED_TITLES:
            title += "_"
        if title.lower() in seen:
            raise BadSpec(f"시트 이름이 겹칩니다: {title}")
        seen.add(title.lower())
        rows = sheet.get("rows")
        if not isinstance(rows, list):
            raise BadSpec("rows 는 행(배열)의 배열이어야 합니다.")
        if len(rows) > MAX_ROWS:
            raise BadSpec(f"한 시트는 {MAX_ROWS}행까지입니다.")
        ws = wb.create_sheet(title)
        for r, row in enumerate(rows, start=1):
            if not isinstance(row, list):
                raise BadSpec(f"{title} {r}행이 배열이 아닙니다.")
            if len(row) > MAX_COLS:
                raise BadSpec(f"한 행은 {MAX_COLS}열까지입니다.")
            cells += len(row)
            if cells > MAX_CELLS:
                raise BadSpec(f"셀은 전부 {MAX_CELLS}개까지입니다.")
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=_value(value, title, r, c))
                if r == 1:
                    cell.font = Font(bold=True)  # 첫 행 = 머리글
    if cells == 0:
        raise BadSpec("빈 표입니다 — 행을 넣어 주세요.")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _value(value: Any, title: str, r: int, c: int) -> Any:
    where = f"{title} {r}행 {c}열"
    if value is None or isinstance(value, bool | int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise BadSpec(f"{where} 값이 수가 아닙니다(무한대·NaN).")
        return value
    if isinstance(value, str):
        if len(value) > MAX_CELL_CHARS:
            raise BadSpec(f"{where} 글이 {MAX_CELL_CHARS}자를 넘습니다.")
        if value.startswith("="):
            _check_formula(value, where)
            return value
        return "".join(ch for ch in value if _xml_ok(ch))
    raise BadSpec(f"{where} 값은 문자열·숫자·불리언·null 만 됩니다.")


def _check_formula(value: str, where: str) -> None:
    if any(ch in _FORMULA_BAD_CHARS for ch in value) or not all(_xml_ok(ch) for ch in value):
        raise BadSpec(
            f"{where} 수식에는 문자열·외부 참조·다른 시트를 쓸 수 없습니다(셀 참조와 숫자만)."
        )
    bad = {f.upper() for f in _FUNC_RE.findall(value)} - _SAFE_FUNCS
    if bad:
        raise BadSpec(
            f"{where} 수식 함수 {sorted(bad)} 는 쓸 수 없습니다(허용: {sorted(_SAFE_FUNCS)})."
        )
